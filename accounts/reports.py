import io
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from django.utils import timezone
from library.models import Borrowing, Book, Category


def make_header_style():
    """Стиль для заголовков таблиц"""
    return {
        'font': Font(bold=True, size=11, color='FFFFFF'),
        'fill': PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid'),
        'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'border': Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        ),
    }


def make_cell_style():
    """Стиль для обычных ячеек"""
    return {
        'alignment': Alignment(vertical='center', wrap_text=True),
        'border': Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        ),
    }


def auto_width(ws):
    """Автоподбор ширины столбцов"""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)


# ---------------------------------------------------------------------------
# ОТЧЁТ 1: Выдачи за период
# ---------------------------------------------------------------------------
def report_borrowings_by_period(start_date, end_date):
    """
    Возвращает HttpResponse с .xlsx-файлом:
    список выдач за указанный период [start_date, end_date].
    """
    borrowings = Borrowing.objects.select_related(
        'book_instance__book', 'user', 'issued_by'
    ).filter(
        borrowed_date__gte=start_date,
        borrowed_date__lte=end_date
    ).order_by('borrowed_date')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Выдачи за период'

    # Заголовок-шапка
    ws.merge_cells('A1:H1')
    ws['A1'] = f'Отчёт: Выдачи за период с {start_date} по {end_date}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    # Колонки
    columns = [
        '№', 'Дата выдачи', 'Читатель (ФИО)', 'Email читателя',
        'Книга', 'Автор(ы)', 'Библиотекарь', 'Дата возврата'
    ]
    header_style = make_header_style()
    cell_style = make_cell_style()

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        for attr, value in header_style.items():
            setattr(cell, attr, value)

    # Данные
    for i, b in enumerate(borrowings, 1):
        book_title = b.book_instance.book.title if b.book_instance and b.book_instance.book else '—'
        authors = ', '.join(
            str(a) for a in b.book_instance.book.authors.all()
        ) if b.book_instance and b.book_instance.book else '—'
        returned = b.returned_date.strftime('%d.%m.%Y') if b.returned_date else '—'

        row = [
            i,
            b.borrowed_date.strftime('%d.%m.%Y'),
            str(b.user),
            b.user.email or '—',
            book_title,
            authors,
            str(b.issued_by) if b.issued_by else '—',
            returned,
        ]
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=3 + i, column=col_idx, value=val)
            for attr, value in cell_style.items():
                setattr(cell, attr, value)

    auto_width(ws)
    ws.freeze_panes = 'A4'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename=borrowings_{start_date}_{end_date}.xlsx'
    )
    return response


# ---------------------------------------------------------------------------
# ОТЧЁТ 2: Должники (просроченные книги)
# ---------------------------------------------------------------------------
def report_overdue():
    """
    Читатели, у которых есть активные выдачи с due_date < сегодня.
    """
    today = timezone.now().date()
    overdue = Borrowing.objects.select_related(
        'book_instance__book', 'user'
    ).filter(
        status='active',
        due_date__lt=today
    ).order_by('due_date')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Должники'

    ws.merge_cells('A1:G1')
    ws['A1'] = f'Отчёт: Должники на {today.strftime("%d.%m.%Y")}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    columns = [
        '№', 'Читатель (ФИО)', 'Email', 'Телефон',
        'Книга', 'Дата выдачи', 'Просрочено (дней)'
    ]
    header_style = make_header_style()
    cell_style = make_cell_style()

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        for attr, value in header_style.items():
            setattr(cell, attr, value)

    for i, b in enumerate(overdue, 1):
        book_title = b.book_instance.book.title if b.book_instance and b.book_instance.book else '—'
        days_overdue = (today - b.due_date).days
        row = [
            i,
            str(b.user),
            b.user.email or '—',
            b.user.phone or '—',
            book_title,
            b.borrowed_date.strftime('%d.%m.%Y'),
            days_overdue,
        ]
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=3 + i, column=col_idx, value=val)
            for attr, value in cell_style.items():
                setattr(cell, attr, value)

    auto_width(ws)
    ws.freeze_panes = 'A4'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename=overdue_{today}.xlsx'
    )
    return response


# ---------------------------------------------------------------------------
# ОТЧЁТ 3: Статистика по категориям
# ---------------------------------------------------------------------------
def report_category_stats():
    """
    Сколько книг в каждой категории и сколько выдач (всего).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Статистика по категориям'

    ws.merge_cells('A1:D1')
    ws['A1'] = 'Отчёт: Статистика по категориям'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    columns = ['№', 'Категория', 'Книг в каталоге', 'Всего выдач']
    header_style = make_header_style()
    cell_style = make_cell_style()

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=3, column=col_idx, value=col_name)
        for attr, value in header_style.items():
            setattr(cell, attr, value)

    categories = Category.objects.all()
    for i, cat in enumerate(categories, 1):
        book_count = Book.objects.filter(category=cat).count()
        borrowing_count = Borrowing.objects.filter(
            book_instance__book__category=cat
        ).count()
        row = [i, cat.name, book_count, borrowing_count]
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=3 + i, column=col_idx, value=val)
            for attr, value in cell_style.items():
                setattr(cell, attr, value)

    auto_width(ws)
    ws.freeze_panes = 'A4'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=category_stats.xlsx'
    return response